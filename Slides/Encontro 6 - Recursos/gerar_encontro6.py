"""Gera o encontro 6 na ordem didática de conteudo_e6.py e atualiza somente sua seção.
Requer python-pptx e python-docx. Execute a partir de qualquer diretório.
"""
from pathlib import Path
from PIL import ImageFont
from copy import deepcopy
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from docx import Document
from docx.shared import Pt as DPt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
LAB=ROOT/'Práticas/encontro-6-observabilidade'
from conteudo_e6 import S, REFS

# Identidade visual: navy, azul e verde; os diagramas são formas vetoriais editáveis.
NAVY='1B2A4A'; BLUE='2A5C8A'; TEAL='007F86'; GRAY='F0F4F8'; INK='26364A'; WHITE='FFFFFF'; ORANGE='A95A17'
def box(sl,x,y,w,h,fill=GRAY,line=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.fill.background()
 return sh
def txt(sl,x,y,w,h,text,size=20,color=INK,bold=False,font='Aptos'):
 sh=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=sh.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE; tf.vertical_anchor=MSO_ANCHOR.TOP; tf._txBody.bodyPr.set("anchorCtr","0")
 tf.margin_left=Inches(.03);tf.margin_right=Inches(.03);tf.margin_top=Inches(.02);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=PP_ALIGN.LEFT;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(9)
 return sh
def label(sl,x,y,w,h,title,body='',color=BLUE):
 box(sl,x,y,w,h)
 font=ImageFont.truetype('DejaVuSans-Bold.ttf',100)
 title_size=min(21 if h>=2.5 else 18,(w-.45)*72*100/font.getlength(title))
 title_shape=txt(sl,x+.17,y+.13,w-.34,.48,title,title_size,color,True)
 title_shape.text_frame.word_wrap=False
 if body:
  compact=h<2.5
  sh=txt(sl,x+.17,y+(.61 if compact else .72),w-.34,h-(.7 if compact else .83),body,15 if compact else 18)
  for pp in sh.text_frame.paragraphs:
   pp.space_after=Pt(3 if compact else 6)
   pp.line_spacing=1.0 if compact else 1.05
def arrow(sl,x,y,w=.48,h=.28):
 sh=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(TEAL);sh.line.fill.background()
def diagram(sl,d):
 if d['kind']=='flow':
  nodes=d['nodes'];w=(12-.45*(len(nodes)-1))/len(nodes)
  for j,(title,body) in enumerate(nodes):
   x=.7+j*(w+.45);label(sl,x,2.2,w,2.8,title,body)
   if j<len(nodes)-1:arrow(sl,x+w+.055,3.4,.34)
  text_block(sl,.85,5.42,11.6,.48,'Leia o fluxo: cada etapa tem um objeto, uma decisão e uma evidência.',18,TEAL,True)
 elif d['kind']=='latency':
  for j,(count,ms,color) in enumerate([(95,20,BLUE),(4,200,TEAL),(1,2000,ORANGE)]):
   yy=2.12+j*1.02
   text_block(sl,.85,yy,3,.65,f'{count} chamadas × {ms} ms',19,INK,True)
   box(sl,4,yy+.05,count/95*5,.5,color)
  text_block(sl,.85,5.4,11.4,.5,'Comprimento das barras = quantidade de chamadas (total: 100).',17,BLUE)
  label(sl,9.35,2.02,3.2,3.3,'RESULTADO','Média: 47 ms\np95: 20 ms\np99: 200 ms\nMáximo: 2.000 ms')
 elif d['kind']=='trace':
  text_block(sl,.85,1.98,11.5,.55,'Uma requisição • eixo horizontal: tempo desde o início',20,BLUE,True)
  box(sl,1.05,2.95,10.8,.65,BLUE)
  text_block(sl,1.2,3.03,10,.5,'Span pai: API • 120 ms',21,WHITE,True)
  box(sl,2.85,4.02,7.2,.65,TEAL)
  text_block(sl,3,4.1,6.8,.5,'Span filho: consulta ao banco • 80 ms',20,WHITE,True)
  for x,t in [(1.05,'0 ms'),(2.85,'20 ms'),(10.05,'100 ms'),(11.65,'120 ms')]:
   text_block(sl,x-.18,2.56,1.4,.35,t,14,INK)
  text_block(sl,.85,5.26,11.5,.65,'Os 80 ms estão dentro dos 120 ms: não somar pai + filho.',21,TEAL,True)
 elif d['kind']=='budget':
  label(sl,.85,2.02,3.35,3.5,'CONTRATO','1.000 operações\nSLO: 99% boas\n10 falhas permitidas')
  text_block(sl,4.65,2.16,7.4,.65,'8 falhas observadas → 80% do orçamento usado',22,INK,True)
  box(sl,4.65,3.24,5.92,.8,ORANGE);box(sl,10.57,3.24,1.48,.8,TEAL)
  text_block(sl,4.8,3.39,5.5,.5,'8 usadas',21,WHITE,True)
  text_block(sl,10.64,3.42,1.3,.4,'2 livres',17,WHITE,True)
  text_block(sl,4.65,4.56,7.4,.85,'992 / 1.000 = 99,2% boas\nObjetivo cumprido nesta janela.',22,BLUE,True)
 elif d['kind']=='dashboard':
  items=[('TRÁFEGO','Chamadas por segundo\nVolume na janela escolhida'),('ERROS','% de operações não boas\nContrato e denominador explícitos'),('LATÊNCIA','p95 em milissegundos\nPopulação e janela definidas'),('SATURAÇÃO','Uso / limite do recurso\nEx.: conexões usadas / disponíveis')]
  for j,(a,b) in enumerate(items):label(sl,.85+(j%2)*6,2.02+(j//2)*1.87,5.7,1.73,a,b)
 elif d['kind']=='timeline':
  for j,(a,b) in enumerate([('10:00 | Impacto','Operação deixa de atender.\nInício da janela de impacto.'),('10:02 | Detecção','Sinais indicam o problema.\n2 min após o início.'),('10:07 | Recuperação','Operação validada novamente.\n7 min de impacto total.')]):
   label(sl,.8+j*4.1,2.2,3.8,2.6,a,b)
   if j<2:arrow(sl,4.62+j*4.1,3.3,.25)
  text_block(sl,.85,5.2,11.8,.7,'Detecção → recuperação: 5 min. Um incidente não é uma média (MTTR).',20,TEAL,True)
 else:raise ValueError(d['kind'])

def text_block(sl,x,y,w,h,value,size=18,color=INK,bold=False,font='Aptos'):
 sh=txt(sl,x,y,w,h,value,size,color,bold,font)
 for pp in sh.text_frame.paragraphs:
  pp.space_after=Pt(5)
  pp.line_spacing=1.05
 return sh

def draw_table(sl,d):
 x=.7;y=1.93;total=12;header=.62;rowh=min(.98,3.4/len(d['rows']))
 for col,(heading,fraction) in enumerate(zip(d['headers'],d['widths'])):
  w=total*fraction
  box(sl,x,y,w-.04,header,NAVY);text_block(sl,x+.13,y+.12,w-.3,.42,heading,18,WHITE,True)
  for i,row in enumerate(d['rows']):
   yy=y+header+.06+i*rowh
   box(sl,x,yy,w-.04,rowh-.045,GRAY if i%2==0 else 'E6F1F3')
   value=row[col];size=16 if len(value)>65 else 17
   cell=text_block(sl,x+.13,yy+.055,w-.3,rowh-.08,value,size,INK,bold=(col==0 and fraction<.35))
   for pp in cell.text_frame.paragraphs:pp.space_after=Pt(0);pp.line_spacing=1.0
  x+=w

def render(p,data):
 made=[]
 for d in data:
  n=d['source_id']
  sl=p.slides.add_slide(min(p.slide_layouts,key=lambda l:len(l.placeholders)))
  sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
  for inherited in list(sl.shapes):inherited._element.getparent().remove(inherited._element)
  font=ImageFont.truetype('DejaVuSans-Bold.ttf',300)
  title_size=min(30,8350/font.getlength(d['title'])*30)
  sh=text_block(sl,.65,.28,12,.65,d['title'],title_size,NAVY,True);sh.text_frame.word_wrap=False
  text_block(sl,.68,1.01,12,.6,d['sub'],16,BLUE)
  box(sl,.68,1.65,1.3,.05,TEAL)
  if d['kind']=='table':draw_table(sl,d)
  elif d['kind']=='code':
   box(sl,.7,1.92,7.55,4.05,NAVY)
   lines=d['code'].strip().splitlines();size=min(18,500/(max(map(len,lines))*.62),245/(len(lines)*1.2))
   sh=text_block(sl,.88,2.09,7.19,3.72,d['code'].strip(),size,WHITE,font='DejaVu Sans Mono');sh.text_frame.word_wrap=False
   for pp in sh.text_frame.paragraphs:pp.space_after=Pt(1)
   for j,(aa,bb) in enumerate(d['cards']):
    step=4.05/len(d['cards']);label(sl,8.5,1.92+j*step,4.1,step-.12,aa,bb)
  elif d['kind']=='cards':
   count=len(d['cards']);w=(12-.28*(count-1))/count
   for j,(aa,bb) in enumerate(d['cards']):label(sl,.7+j*(w+.28),2.02,w,3.65,aa,bb)
  else:diagram(sl,d)
  box(sl,.7,6.13,12,.78,'E6F1F3')
  text_block(sl,.88,6.24,11.64,.59,d['takeaway'],17 if len(d['takeaway'])>135 else 18,TEAL,True)
  footer='UNIFOR • DevOps • Encontro 6  |  '+d['section']
  if d.get('study_only'):footer+='  |  CONSULTA / APROFUNDAMENTO'
  text_block(sl,.7,7.1,10.8,.22,footer,10,BLUE)
  text_block(sl,11.8,7.04,.85,.3,f'E6·{n:02}',11,BLUE)
  notes=f"E6.{n:02} — {d['title']}\n\nCOMO EXPLICAR\n{d['note']}\n\nCONCLUSÃO\n{d['takeaway']}"
  if d['question']:notes+='\n\nPERGUNTAS\n'+d['question']
  if d['answer']:notes+='\n\nRESPOSTAS\n'+d['answer']
  sl.notes_slide.notes_text_frame.text=notes+'\n\nREFERÊNCIAS\n'+'\n'.join(t+': '+u for t,u in REFS)
  made.append(sl)
 return made
def slide_heading(sl):
 return next((sh.text.strip() for sh in sl.shapes if sh.has_text_frame and sh.text.strip()),'')

def add_cover(p,reference):
 """Reutiliza a capa de outro encontro, incluindo layout, fundo e imagens."""
 sl=p.slides.add_slide(reference.slide_layout)
 for sh in list(sl.shapes):sh._element.getparent().remove(sh._element)
 if reference._element.cSld.bg is not None:
  old=sl._element.cSld.bg
  if old is not None:sl._element.cSld.remove(old)
  sl._element.cSld.insert(0,deepcopy(reference._element.cSld.bg))
 for sh in reference.shapes:
  el=deepcopy(sh._element)
  for blip in el.xpath('.//a:blip'):
   key='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
   rid=blip.get(key)
   if rid:
    blob=reference.part.related_part(rid).blob
    _,newrid=sl.part.get_or_add_image_part(BytesIO(blob))
    blip.set(key,newrid)
  sl.shapes._spTree.insert_element_before(el,'p:extLst')
 for sh in sl.shapes:
  if not sh.has_text_frame:continue
  for par in sh.text_frame.paragraphs:
   for run in par.runs:
    value=run.text
    if value.strip().isdigit():value='E6'
    run.text=value
 sl.notes_slide.notes_text_frame.text='Capa do Encontro 6 — 26/09/2026. Observabilidade, IA aplicada e projeto final. Avance para o conteúdo E6·01; as referências E6·NN permanecem válidas.'
 return sl

def export_encontro(main,target):
 """Preserva o tema institucional também na apresentação independente."""
 p=Presentation(main)
 start=next(i for i,sl in enumerate(p.slides) if slide_heading(sl)=='Encontro 6')
 end=next(i for i,sl in enumerate(p.slides) if i>start and slide_heading(sl)=='Anexos')
 ids=p.slides._sldIdLst
 for i,el in reversed(list(enumerate(list(ids)))):
  if not start<=i<end:p.part.drop_rel(el.rId);ids.remove(el)
 p.save(target)

def replace_section(path,data):
 p=Presentation(path);ids=p.slides._sldIdLst
 def marker(s,text):
  first=next((sh.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text.strip()),'')
  return first==text or first.startswith(text+' |')
 start=next(i for i,s in enumerate(p.slides) if marker(s,'Encontro 6'))
 end=next(i for i,s in enumerate(p.slides) if i>start and (marker(s,'Anexos') or marker(s,'Apêndice')))
 old=list(ids)[start:end];cover=add_cover(p,p.slides[start]);new=[cover]+render(p,data);newids=list(ids)[-len(new):]
 for e in old:p.part.drop_rel(e.rId);ids.remove(e)
 for e in newids:ids.remove(e)
 for j,e in enumerate(newids):ids.insert(start+j,e)
 p.save(path)
 return start+1,start+len(new)
if __name__=='__main__':
 main=ROOT/'Slides/Desenvolvimento de Software Integrado - DevOps - Slides.pptx'
 interval=replace_section(main,S)
 export_encontro(main,OUT/'Encontro 6 - Observabilidade e Projeto Final.pptx')

 replace_section(ROOT/'Slides/Desenvolvimento de Software Integrado - DevOps - Guia de Aula.pptx',[d for d in S if not d['study_only']])
 path=ROOT/'Slides/Desenvolvimento de Software Integrado - DevOps - Guia de Aula.docx';doc=Document(path);body=doc._element.body
 begin=next(p for p in doc.paragraphs if p.text.startswith('Parte 8 — Encontro 6'))
 end=next(p for p in doc.paragraphs if p.text.startswith('Parte 9 —'))
 elements=list(body);a=elements.index(begin._p);b=elements.index(end._p)
 for el in elements[a:b]:body.remove(el)
 def insert_para(text,style=None):
  pp=doc.add_paragraph(text);end._p.addprevious(pp._p)
  if style:
   pp.style=next(ss for ss in doc.styles if ss.name==style);pp.paragraph_format.keep_with_next=True
   for run in pp.runs:run.font.bold=True;run.font.size=DPt(18 if style=='Heading 1' else 13)
  return pp
 intro=f'{len(S)} slides de conteúdo + capa institucional • 26/09/2026 • 4h presenciais. No conjunto principal: posições {interval[0]}–{interval[1]}. Revisão: 23/09/2026.'
 usage='Sequência: fundamentos e sinais → indicadores e objetivos → investigação → demonstração → revisão de IA → preparação e apresentação do projeto. Preservar 35 minutos de preparação e 56 de apresentações (até sete equipes, falas de 5–7 minutos). O exemplo da API não substitui o projeto escolhido. As verificações têm gabaritos comentados; os últimos quatro slides são consulta.'
 insert_para('Parte 8 — Encontro 6: Observabilidade, IA Aplicada e Projeto Final','Heading 1').paragraph_format.page_break_before=True
 insert_para(intro);insert_para(usage)
 md=['# Encontro 6 — Guia ampliado do professor','',intro,'',usage,'']
 for d in S:
  i=d['source_id'];title=f'E6.{i:02} — {d["title"]}'
  insert_para(title,'Heading 2');insert_para(d['sub'])
  md += ['## '+title,'',d['sub'],'']
  for heading,content in d['cards']:
   insert_para(heading+': '+content.replace('\n',' '));md += ['**'+heading+':** '+content.replace('\n',' '),'']
  if d['kind']=='flow':
   for heading,content in d['nodes']:
    insert_para(heading+': '+content.replace('\n',' '));md += ['**'+heading+':** '+content.replace('\n',' '),'']
  if d['kind']=='table':
   table=doc.add_table(rows=1,cols=len(d['headers']));end._p.addprevious(table._tbl)
   try:table.style='Table Grid'
   except KeyError:pass
   for cell,heading in zip(table.rows[0].cells,d['headers']):cell.text=heading
   for values in d['rows']:
    for cell,value in zip(table.add_row().cells,values):cell.text=value
   for row in table.rows:
    for cell in row.cells:
     for pp in cell.paragraphs:
      for run in pp.runs:run.font.size=DPt(10)
   md += ['| '+' | '.join(d['headers'])+' |','| '+' | '.join('---' for _ in d['headers'])+' |']
   md += ['| '+' | '.join(v.replace('|',r'\|') for v in row)+' |' for row in d['rows']]+['']
  if d['code']:
   pp=insert_para(d['code'].strip())
   for run in pp.runs:run.font.name='Consolas';run.font.size=DPt(9)
   md += ['```',d['code'].strip(),'```','']
  insert_para('Condução e explicação: '+d['note']);insert_para('Conclusão prática: '+d['takeaway'])
  md += ['**Como explicar:** '+d['note'],'','**Conclusão prática:** '+d['takeaway'],'']
 insert_para('Referências — documentação oficial','Heading 2')
 for title,url in REFS:insert_para(title+' — '+url)
 for existing in list(doc.settings.element.findall(qn('w:updateFields'))):doc.settings.element.remove(existing)
 update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
 doc.save(path)
 md+=['## Referências','']+[f'- [{t}]({u})' for t,u in REFS]
 (OUT/'Guia ampliado do professor.md').write_text('\n'.join(md)+'\n')
 (OUT/'conteudo.json').write_text(json.dumps(S,ensure_ascii=False,indent=2)+'\n')
 idx=['# Verificações e respostas comentadas','','As cinco verificações aparecem depois da explicação e da demonstração de seus blocos. Cada uma tem três situações, seguidas de um slide com respostas explicadas. Os casos resolvidos mostram sinais, cálculos, investigação, recuperação e revisão de IA.','','Use o rodapé E6·NN. A página no PDF independente corresponde ao número + 1, devido à capa.','','| Questões | Gabarito | Bloco |','|---|---|---|']
 for d in S:
  if d['role']=='resposta':
   q=next(x for x in S if x['id']==d['answers_for'])
   idx.append(f"| E6·{q['source_id']:02} | E6·{d['source_id']:02} | {q['title']} |")
 for d in S:
  if d['role']=='resposta':
   q=next(x for x in S if x['id']==d['answers_for']);idx+=['','## '+q['title'],'']
   for question,answer in zip(q['questions'],d['answers']):idx+=['**Pergunta:** '+question,'','**Resposta explicada:** '+answer,'']
 idx+=['## Atividade no projeto','','A equipe adapta instrumentação, relato do incidente e roteiro de apresentação ao próprio projeto. Os exemplos preenchidos estão na pasta Práticas/encontro-6-observabilidade. Distinguir sempre o que foi planejado, executado e validado.']
 (OUT/'Índice de perguntas e respostas.md').write_text('\n'.join(idx)+'\n')
 print(f'{len(S)} slides E6 + capa; intervalo principal {interval}; guias atualizados.')
