"""Gera o encontro 4 na ordem didática de conteudo_e4.py e atualiza somente sua seção.
Requer python-pptx e python-docx. Execute a partir de qualquer diretório.
"""
from pathlib import Path
from PIL import ImageFont
from copy import deepcopy
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE, MSO_ANCHOR
from docx import Document
from docx.shared import Pt as DPt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json
ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
LAB=ROOT/'Práticas/encontro-4-containers'
from conteudo_e4 import S, REFS

# Identidade visual: navy, azul e verde; os diagramas são formas vetoriais editáveis.
NAVY='1B2A4A'; BLUE='2A5C8A'; TEAL='007F86'; GRAY='F0F4F8'; INK='26364A'; WHITE='FFFFFF'; ORANGE='A95A17'
def box(sl,x,y,w,h,fill=GRAY,line=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor.from_string(fill); sh.line.fill.background()
 return sh
def txt(sl,x,y,w,h,text,size=20,color=INK,bold=False,font='Aptos'):
 sh=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=sh.text_frame; tf.word_wrap=True; tf.auto_size=MSO_AUTO_SIZE.NONE; tf.vertical_anchor=MSO_ANCHOR.TOP; tf._txBody.bodyPr.set("anchorCtr","0")
 tf.margin_left=Inches(.03);tf.margin_right=Inches(.03);tf.margin_top=Inches(.02);tf.margin_bottom=0
 for i,line in enumerate(text.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=PP_ALIGN.LEFT;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(9)
 return sh
def label(sl,x,y,w,h,title,body='',color=BLUE):
 box(sl,x,y,w,h)
 font=ImageFont.truetype('DejaVuSans-Bold.ttf',100)
 title_size=min(21 if h>=2.5 else 18,(w-.45)*72*100/font.getlength(title))
 title_shape=txt(sl,x+.17,y+.13,w-.34,.48,title,title_size,color,True)
 title_shape.text_frame.word_wrap=False
 if body:
  compact=h<2.5
  sh=txt(sl,x+.17,y+(.61 if compact else .72),w-.34,h-(.7 if compact else .83),body,15 if compact else 18)
  for pp in sh.text_frame.paragraphs:
   pp.space_after=Pt(3 if compact else 6)
   pp.line_spacing=1.0 if compact else 1.05
def arrow(sl,x,y,w=.48,h=.28):
 sh=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(TEAL);sh.line.fill.background()
def diagram(sl,kind):
 if kind=='process':
  label(sl,.7,2.15,3.4,2.5,'ARQUIVO','server.js\nCódigo guardado no disco.')
  arrow(sl,4.27,3.2,.65)
  label(sl,5.15,2.15,3.1,2.5,'COMANDO','node server.js\nNode inicia a execução.')
  arrow(sl,8.42,3.2,.65)
  label(sl,9.3,2.15,3.35,2.5,'PROCESSO','Usa CPU e memória.\nAbre a porta HTTP.')
  text_block(sl,.85,5.05,11.8,.6,'O arquivo permanece; o processo pode iniciar e terminar.',22,TEAL,True)
 elif kind=='instances':
  label(sl,.75,2.7,3.1,2.25,'IMAGEM','Node + bibliotecas\nCódigo + comando padrão')
  arrow(sl,4.05,3.6,.65)
  label(sl,5,1.95,7.6,1.6,'CONTAINER A','Mesma origem • configuração A • camada gravável A')
  label(sl,5,3.85,7.6,1.6,'CONTAINER B','Mesma origem • configuração B • camada gravável B')
 elif kind=='runtimes':
  label(sl,.75,2,5.55,3.5,'RUNTIME DE CONTAINERS','containerd gerencia o ciclo de vida.\nrunc prepara e inicia o processo isolado.\nAtuam na execução do container.')
  label(sl,6.65,2,5.95,3.5,'RUNTIME DA APLICAÇÃO','Node.js executa server.js.\nA API usa Node dentro da instância.\nAtua na execução do código.')
 elif kind=='engine':
  for j,(t,b) in enumerate([('CLI','docker run\nPedido do usuário'),('ENGINE','Gerencia objetos\ne coordena a execução'),('RUNTIME','containerd / runc\nPrepara e inicia'),('PROCESSO','Node + API\nAtende requisições')]):
   label(sl,.7+j*3.1,2.1,2.7,2.65,t,b)
   if j<3:arrow(sl,3.46+j*3.1,3.3,.3)
  box(sl,3.8,5.05,5.8,.62,'E5F1F3');text_block(sl,3.94,5.16,5.5,.4,'Registry → Engine: imagem, se necessária',16,TEAL,True)
 elif kind=='compose_model':
  for j,(t,b) in enumerate([('compose.yaml','Declara serviços,\nredes e volumes.'),('COMPOSE','Lê a configuração\ne solicita operações.'),('ENGINE','Cria e gerencia\nos recursos.')]):
   label(sl,.7+j*4.1,2,3.7,2.2,t,b)
   if j<2:arrow(sl,4.46+j*4.1,2.9,.3)
  box(sl,1,4.75,11.3,.8,NAVY);text_block(sl,1.2,4.94,10.9,.5,'Resultado do exemplo: app + db + rede compartilhada + volume dados',21,WHITE,True)
 elif kind=='project':
  for j,(t,b) in enumerate([('E1–E3 | BASE','Projeto escolhido\nColaboração e código\nTestes e pipeline CI'),('E4 | EXECUÇÃO','Imagem da aplicação\nServiços integrados\nDiagnóstico com evidência'),('E5–E6 | ENTREGA','Configuração e promoção\nRollback e governança\nDemonstração final')]):
   label(sl,.7+j*4.1,2.1,3.7,3.4,t,b)
   if j<2:arrow(sl,4.46+j*4.1,3.5,.3)
 elif kind=='vm':
  for x,title in [(.7,'MÁQUINAS VIRTUAIS'),(6.9,'CONTAINERS LINUX')]:
   text_block(sl,x,1.87,5.7,.4,title,22,BLUE,True)
   for j in range(2):
    xx=x+j*2.9;box(sl,xx,2.45,2.7,1.65);text_block(sl,xx+.14,2.62,2.4,1.3,'Aplicação + libs\n'+('SO + kernel convidado' if x<1 else 'Processo isolado'),19)
   for yy,st in [(4.32,'Hypervisor' if x<1 else 'Runtime de containers'),(5.12,'Host / hardware' if x<1 else 'Kernel do host Linux')]:
    box(sl,x,yy,5.6,.63,BLUE if yy<5 else NAVY);text_block(sl,x+.18,yy+.11,5.2,.45,st,21,WHITE)
 elif kind=='network':
  label(sl,.7,2.6,3,2.1,'CLIENTE NO HOST','localhost:8080\nConexão HTTP');arrow(sl,3.85,3.4,.6)
  box(sl,4.65,1.95,7.95,3.7,'E5F1F3');text_block(sl,4.9,2.1,7.4,.4,'REDE DO COMPOSE',19,TEAL,True)
  label(sl,4.95,2.75,3,2,'app','0.0.0.0:3000');arrow(sl,8.08,3.5,.5)
  label(sl,8.75,2.75,3.5,2,'db','PostgreSQL :5432')
  text_block(sl,8.85,4.96,3.5,.4,'SQL • volume: dados',18,TEAL,True)
 elif kind=='cache':
  for x,title,steps in [(.7,'Ordem que desperdiça cache',['COPY . .','RUN npm ci','Código mudou → reinstala']),(6.9,'Ordem do laboratório',['COPY package*.json ./','RUN npm ci','COPY server.js ./'])]:
   text_block(sl,x,1.9,5.6,.5,title,22,BLUE,True)
   for j,v in enumerate(steps):
    box(sl,x,2.6+j*.98,5.6,.76);text_block(sl,x+.15,2.76+j*.98,5.3,.48,v,18,INK,font='DejaVu Sans Mono')
 elif kind=='storage':
  for j,(t,b) in enumerate([('Camada gravável','Arquivos da instância.\nRemover o container elimina essa camada.'),('Volume nomeado','Gerenciado pelo Docker.\nO recurso pode persistir após a recriação.'),('Bind mount','Caminho do host.\nNo exemplo: init.sql, montado somente leitura.')]):label(sl,.7+j*4.1,2.05,3.8,3.35,t,b)
 elif kind=='health':
  for j,(t,b) in enumerate([('RUNNING','Processo em execução.\nNão prova acesso ao banco.'),('HEALTHY','O check configurado passou.\nSeu alcance depende do teste.'),('FUNCIONAL','A operação do usuário funciona.\nEx.: gravar e ler /visitas.')]):label(sl,.7+j*4.1,2.15,3.8,3.35,t,b)
 elif kind=='trouble':
  for j,(t,b) in enumerate([('1 • Fase','Build ou execução?'),('2 • Estado','ps -a'),('3 • Evidência','logs'),('4 • Integração','Rede e banco'),('5 • Prova','Requisição real')]):label(sl,.65+j*2.55,2.5,2.35,2.65,t,b)
 else:raise ValueError('Diagrama desconhecido: '+kind)

def text_block(sl,x,y,w,h,value,size=18,color=INK,bold=False,font='Aptos'):
 sh=txt(sl,x,y,w,h,value,size,color,bold,font)
 for pp in sh.text_frame.paragraphs:
  pp.space_after=Pt(5)
  pp.line_spacing=1.05
 return sh

def draw_table(sl,d):
 x=.7;y=1.93;total=12;header=.62;rowh=min(.98,3.4/len(d['rows']))
 for col,(heading,fraction) in enumerate(zip(d['headers'],d['widths'])):
  w=total*fraction
  box(sl,x,y,w-.04,header,NAVY);text_block(sl,x+.13,y+.12,w-.3,.42,heading,18,WHITE,True)
  for i,row in enumerate(d['rows']):
   yy=y+header+.06+i*rowh
   box(sl,x,yy,w-.04,rowh-.045,GRAY if i%2==0 else 'E5F1F3')
   value=row[col];size=16 if len(value)>65 else 17
   cell=text_block(sl,x+.13,yy+.055,w-.3,rowh-.08,value,size,INK,bold=(col==0 and fraction<.35))
   for pp in cell.text_frame.paragraphs:pp.space_after=Pt(0);pp.line_spacing=1.0
  x+=w

def render(p,data):
 made=[]
 for d in data:
  n=d['source_id']
  sl=p.slides.add_slide(min(p.slide_layouts,key=lambda l:len(l.placeholders)))
  sl.background.fill.solid();sl.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
  for inherited in list(sl.shapes):inherited._element.getparent().remove(inherited._element)
  font=ImageFont.truetype('DejaVuSans-Bold.ttf',300)
  title_size=min(30,8350/font.getlength(d['title'])*30)
  sh=text_block(sl,.65,.28,12,.65,d['title'],title_size,NAVY,True);sh.text_frame.word_wrap=False
  text_block(sl,.68,1.01,12,.6,d['sub'],16,BLUE)
  box(sl,.68,1.65,1.3,.05,TEAL)
  if d['kind']=='table':draw_table(sl,d)
  elif d['kind']=='code':
   box(sl,.7,1.92,7.55,4.05,NAVY)
   lines=d['code'].strip().splitlines();size=min(18,500/(max(map(len,lines))*.62),245/(len(lines)*1.2))
   sh=text_block(sl,.88,2.09,7.19,3.72,d['code'].strip(),size,WHITE,font='DejaVu Sans Mono');sh.text_frame.word_wrap=False
   for pp in sh.text_frame.paragraphs:pp.space_after=Pt(1)
   for j,(aa,bb) in enumerate(d['cards']):
    step=4.05/len(d['cards']);label(sl,8.5,1.92+j*step,4.1,step-.12,aa,bb)
  elif d['kind']=='cards':
   count=len(d['cards']);w=(12-.28*(count-1))/count
   for j,(aa,bb) in enumerate(d['cards']):label(sl,.7+j*(w+.28),2.02,w,3.65,aa,bb)
  else:diagram(sl,d['kind'])
  box(sl,.7,6.13,12,.78,'E5F1F3')
  text_block(sl,.88,6.24,11.64,.59,d['takeaway'],17 if len(d['takeaway'])>135 else 18,TEAL,True)
  footer='UNIFOR • DevOps • Encontro 4  |  '+d['section']
  if d.get('study_only'):footer+='  |  CONSULTA / APROFUNDAMENTO'
  text_block(sl,.7,7.1,10.8,.22,footer,10,BLUE)
  text_block(sl,11.8,7.04,.85,.3,f'E4·{n:02}',11,BLUE)
  notes=f"E4.{n:02} — {d['title']}\n\nCOMO EXPLICAR\n{d['note']}\n\nCONCLUSÃO\n{d['takeaway']}"
  if d['question']:notes+='\n\nPERGUNTAS\n'+d['question']
  if d['answer']:notes+='\n\nRESPOSTAS\n'+d['answer']
  sl.notes_slide.notes_text_frame.text=notes+'\n\nREFERÊNCIAS\n'+'\n'.join(t+': '+u for t,u in REFS)
  made.append(sl)
 return made
def slide_heading(sl):
 return next((sh.text.strip() for sh in sl.shapes if sh.has_text_frame and sh.text.strip()),'')

def add_cover(p,reference):
 """Reutiliza a capa de outro encontro, incluindo layout, fundo e imagens."""
 sl=p.slides.add_slide(reference.slide_layout)
 for sh in list(sl.shapes):sh._element.getparent().remove(sh._element)
 if reference._element.cSld.bg is not None:
  old=sl._element.cSld.bg
  if old is not None:sl._element.cSld.remove(old)
  sl._element.cSld.insert(0,deepcopy(reference._element.cSld.bg))
 for sh in reference.shapes:
  el=deepcopy(sh._element)
  for blip in el.xpath('.//a:blip'):
   key='{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed'
   rid=blip.get(key)
   if rid:
    blob=reference.part.related_part(rid).blob
    if sh.shape_type==13 and 3<sh.left/914400<9 and .8<sh.top/914400<1.8:
     blob=(OUT/'icone-containers.png').read_bytes()
    _,newrid=sl.part.get_or_add_image_part(BytesIO(blob))
    blip.set(key,newrid)
  sl.shapes._spTree.insert_element_before(el,'p:extLst')
 for sh in sl.shapes:
  if not sh.has_text_frame:continue
  for par in sh.text_frame.paragraphs:
   for run in par.runs:
    value=run.text.replace('Encontro 5','Encontro 4').replace('25/09/2026','24/09/2026')
    value=value.replace('CD, configuração, segurança e governança de IA','Containers, integração e troubleshooting')
    if value.strip().isdigit():value='E4'
    run.text=value
 sl.notes_slide.notes_text_frame.text='Capa do Encontro 4 — 24/09/2026. Containers, integração e troubleshooting. Avance para o conteúdo E4·01; as referências E4·NN permanecem válidas.'
 return sl

def export_encontro(main,target):
 """Preserva o tema institucional também na apresentação independente."""
 p=Presentation(main)
 start=next(i for i,sl in enumerate(p.slides) if slide_heading(sl)=='Encontro 4')
 end=next(i for i,sl in enumerate(p.slides) if i>start and slide_heading(sl)=='Encontro 5')
 ids=p.slides._sldIdLst
 for i,el in reversed(list(enumerate(list(ids)))):
  if not start<=i<end:p.part.drop_rel(el.rId);ids.remove(el)
 p.save(target)

def replace_section(path,data):
 p=Presentation(path);ids=p.slides._sldIdLst
 def marker(s,text):
  first=next((sh.text.strip() for sh in s.shapes if sh.has_text_frame and sh.text.strip()),'')
  return first==text or first.startswith(text+' |')
 start=next(i for i,s in enumerate(p.slides) if marker(s,'Encontro 4'))
 end=next(i for i,s in enumerate(p.slides) if i>start and marker(s,'Encontro 5'))
 old=list(ids)[start:end];cover=add_cover(p,p.slides[end]);new=[cover]+render(p,data);newids=list(ids)[-len(new):]
 for e in old:p.part.drop_rel(e.rId);ids.remove(e)
 for e in newids:ids.remove(e)
 for j,e in enumerate(newids):ids.insert(start+j,e)
 p.save(path)
 return start+1,start+len(new)
if __name__=='__main__':
 main=ROOT/'Slides/Desenvolvimento de Software Integrado - DevOps - Slides.pptx'
 interval=replace_section(main,S)
 export_encontro(main,OUT/'Encontro 4 - Containers e Compose.pptx')
 chosen={'percurso','agenda','ingredientes','runtime_app','container','runtime_container','docker_flow','compose','vm_comparison','resources_map','demo_version','demo_http','observe','check_base','check_base_resposta','dockerfile','image_commands','check_build','check_build_resposta','network','storage','health','resource_limits','restart_policy','compose_demo','persistence_demo','check_integration','check_integration_resposta','diagnosis','case_localhost','case_bind','case_dependency','project_bridge','project_plan','project_lab','acceptance','check_final','check_final_resposta','next'}
 replace_section(ROOT/'Slides/Desenvolvimento de Software Integrado - DevOps - Guia de Aula.pptx',[d for d in S if d['id'] in chosen])
 path=ROOT/'Slides/Desenvolvimento de Software Integrado - DevOps - Guia de Aula.docx';doc=Document(path);body=doc._element.body
 begin=next(p for p in doc.paragraphs if p.text.startswith('Parte 6 — Encontro 4'))
 end=next(p for p in doc.paragraphs if p.text.startswith('Parte 7 — Encontro 5'))
 elements=list(body);a=elements.index(begin._p);b=elements.index(end._p)
 for el in elements[a:b]:body.remove(el)
 def insert_para(text,style=None):
  pp=doc.add_paragraph(text);end._p.addprevious(pp._p)
  if style:
   pp.style=next(ss for ss in doc.styles if ss.name==style);pp.paragraph_format.keep_with_next=True
   for run in pp.runs:run.font.bold=True;run.font.size=DPt(18 if style=='Heading 1' else 13)
  return pp
 intro=f'{len(S)} slides de conteúdo + capa institucional • 24/09/2026 • 4h presenciais. No conjunto principal: posições {interval[0]}–{interval[1]}. Revisão: 22/09/2026.'
 usage='Sequência: explicar os fundamentos → demonstrar → verificar com gabarito → integrar → diagnosticar → aplicar ao projeto. As verificações aparecem depois das explicações. Os recursos avançados estão identificados como consulta. Preservar 70 minutos de trabalho das equipes. O exemplo da API não substitui o projeto escolhido. Notas do apresentador e o roteiro de quatro horas apoiam a condução.'
 insert_para('Parte 6 — Encontro 4: Containers, Docker Compose e Troubleshooting','Heading 1')
 insert_para(intro);insert_para(usage)
 md=['# Encontro 4 — Guia ampliado do professor','',intro,'',usage,'']
 for d in S:
  i=d['source_id'];title=f'E4.{i:02} — {d["title"]}'
  insert_para(title,'Heading 2');insert_para(d['sub'])
  md += ['## '+title,'',d['sub'],'']
  for heading,content in d['cards']:
   insert_para(heading+': '+content.replace('\n',' '));md += ['**'+heading+':** '+content.replace('\n',' '),'']
  if d['kind']=='table':
   table=doc.add_table(rows=1,cols=len(d['headers']));end._p.addprevious(table._tbl)
   try:table.style='Table Grid'
   except KeyError:pass
   for cell,heading in zip(table.rows[0].cells,d['headers']):cell.text=heading
   for values in d['rows']:
    for cell,value in zip(table.add_row().cells,values):cell.text=value
   for row in table.rows:
    for cell in row.cells:
     for pp in cell.paragraphs:
      for run in pp.runs:run.font.size=DPt(10)
   md += ['| '+' | '.join(d['headers'])+' |','| '+' | '.join('---' for _ in d['headers'])+' |']
   md += ['| '+' | '.join(v.replace('|',r'\|') for v in row)+' |' for row in d['rows']]+['']
  if d['code']:
   pp=insert_para(d['code'].strip())
   for run in pp.runs:run.font.name='Consolas';run.font.size=DPt(9)
   md += ['```',d['code'].strip(),'```','']
  insert_para('Condução e explicação: '+d['note']);insert_para('Conclusão prática: '+d['takeaway'])
  md += ['**Como explicar:** '+d['note'],'','**Conclusão prática:** '+d['takeaway'],'']
 insert_para('Referências — documentação oficial','Heading 2')
 for title,url in REFS:insert_para(title+' — '+url)
 for existing in list(doc.settings.element.findall(qn('w:updateFields'))):doc.settings.element.remove(existing)
 update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');doc.settings.element.append(update)
 doc.save(path)
 md+=['## Referências','']+[f'- [{t}]({u})' for t,u in REFS]
 (OUT/'Guia ampliado do professor.md').write_text('\n'.join(md)+'\n')
 (OUT/'conteudo.json').write_text(json.dumps(S,ensure_ascii=False,indent=2)+'\n')
 idx=['# Verificações e respostas comentadas','','As quatro verificações aparecem depois da explicação e da demonstração de seus blocos. Cada uma tem três situações, seguidas de um slide com respostas explicadas. Os casos de diagnóstico também mostram sintoma, causa, correção e prova.','','Use o rodapé E4·NN. A página no PDF independente corresponde ao número + 1, devido à capa.','','| Questões | Gabarito | Bloco |','|---|---|---|']
 for d in S:
  if d['role']=='resposta':
   q=next(x for x in S if x['id']==d['answers_for'])
   idx.append(f"| E4·{q['source_id']:02} | E4·{d['source_id']:02} | {q['title']} |")
 for d in S:
  if d['role']=='resposta':
   q=next(x for x in S if x['id']==d['answers_for']);idx+=['','## '+q['title'],'']
   for question,answer in zip(q['questions'],d['answers']):idx+=['**Pergunta:** '+question,'','**Resposta explicada:** '+answer,'']
 idx+=['## Atividade no projeto','','O plano técnico pede decisões da própria equipe. Não há uma única stack correta: as decisões devem corresponder ao projeto e à execução. O slide “Traduzir o exemplo para a aplicação da equipe” mostra o mapeamento e “Exemplo preenchido: evidência de uma falha corrigida” oferece um modelo resolvido.']
 (OUT/'Índice de perguntas e respostas.md').write_text('\n'.join(idx)+'\n')
 print(f'{len(S)} slides E4 + capa; intervalo principal {interval}; guias atualizados.')
